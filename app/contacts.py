"""Local contact book and recognition of common school contact lists."""

from dataclasses import dataclass, field
from itertools import islice, zip_longest
from pathlib import Path
import json
import os
import re
import tempfile

from openpyxl import load_workbook
from zipfile import ZipFile

from app.tools.office import check_zip


CONTACT_LIMIT = 1000
FIELDS = {
    'name': '学生姓名', 'father': '父亲姓名', 'fatherPhone': '父亲联系电话',
    'mother': '母亲姓名', 'motherPhone': '母亲联系电话',
    'guardian': '家长 / 监护人1姓名', 'phone': '家长 / 监护人1电话',
    'guardian2': '家长 / 监护人2姓名', 'guardian2Phone': '家长 / 监护人2电话',
    'address': '家庭住址', 'notes': '备注 / 待核对原文',
}
PHONE_FIELDS = {'phone', 'fatherPhone', 'motherPhone', 'guardian2Phone'}
OPTIONAL_OLD_FIELDS = {'guardian', 'guardian2', 'guardian2Phone'}
ROLE_WORDS = r'父亲|爸爸|父方|母亲|妈妈|母方|家长[12一二]?|监护人[12一二]?|第一监护人|第二监护人|联系人[12一二]?|外公|外婆|爷爷|奶奶'
PHONE_WORDS = r'手机号码|联系电话|电话号码|联系方式|联系手机|手机号|电话|手机'
NAME_WORDS = r'姓名|名字'
STUDENT_WORDS = r'(?:学生|孩子|儿童|幼儿|宝宝|学员)(?:姓名|名字)?'
LABELS = re.compile(rf'(?:{ROLE_WORDS})(?:{PHONE_WORDS}|{NAME_WORDS})?|{STUDENT_WORDS}|{PHONE_WORDS}|{NAME_WORDS}|家庭住址|家庭地址|现住址|居住地址|居住地|住址|地址')
PHONE = re.compile(r'(?<!\d)(?:\+?86[\s-]*)?(1[3-9]\d[\s-]*\d{4}[\s-]*\d{4})(?!\d)')
CONTINUATION = re.compile(rf'^(?:(?:{ROLE_WORDS})(?:{PHONE_WORDS}|{NAME_WORDS})?|{PHONE_WORDS}|家庭住址|家庭地址|住址|地址)[:：\s]')


def blank_contact():
    return {key: '' for key in FIELDS}


def normalize(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str('' if value is None else value).translate(str.maketrans('０１２３４５６７８９', '0123456789')).strip()


def extract_phones(value):
    return ' / '.join(dict.fromkeys(re.sub(r'\D', '', match)[-11:] for match in PHONE.findall(normalize(value))))


def role_of(label):
    if re.search(r'父亲|爸爸|父方', label):
        return 'father'
    if re.search(r'母亲|妈妈|母方', label):
        return 'mother'
    if re.search(r'家长|监护人|联系人|外公|外婆|爷爷|奶奶', label):
        return 'guardian2' if re.search(r'2|二|第二', label) else 'guardian'
    return None


def phone_field(role):
    return {'father': 'fatherPhone', 'mother': 'motherPhone',
            'guardian': 'phone', 'guardian2': 'guardian2Phone'}[role]


def header_field(value):
    text = re.sub(r'[\s:：()（）_\-]', '', normalize(value))
    role = role_of(text)
    is_phone = re.search(PHONE_WORDS, text)
    if re.fullmatch(rf'(?:{ROLE_WORDS})(?:{NAME_WORDS}|{PHONE_WORDS})?', text):
        return phone_field(role) if is_phone else role
    if re.fullmatch(rf'(?:{NAME_WORDS}|{PHONE_WORDS})(?:{ROLE_WORDS})', text):
        return phone_field(role) if is_phone else role
    if re.fullmatch(r'家庭住址|家庭地址|现住址|居住地址|居住地|住址|地址|家庭所在地', text):
        return 'address'
    if re.fullmatch(rf'(?:{PHONE_WORDS})[12一二]?', text):
        return 'guardian2Phone' if re.search(r'2|二', text) else 'phone'
    if re.fullmatch(rf'(?:{STUDENT_WORDS}|{NAME_WORDS})', text):
        return 'name'
    return None


def _append(contact, key, value):
    if value:
        contact[key] = ' / '.join(dict.fromkeys(filter(None, [*contact[key].split(' / '), *value.split(' / ')])))


def _without_phones(value):
    return re.sub(r'^[\s:：,，;；|、()（）]+|[\s,，;；|、()（）]+$', '', PHONE.sub(' ', value))


def _valid_name(value):
    return bool(re.fullmatch(r'[\u4e00-\u9fff·]{2,12}|[A-Za-z][A-Za-z .\'-]{1,59}', value))


def _check(contact, problems):
    if not contact['name']:
        problems.append('缺少学生姓名')
    if not any(contact[key] for key in PHONE_FIELDS):
        problems.append('未识别到有效的11位家长手机号')


@dataclass
class Recognition:
    contacts: list[dict] = field(default_factory=list)
    skipped: int = 0
    issues: list[str] = field(default_factory=list)
    sheets: int = 0
    no_headers: int = 0
    has_header: bool = False


def _parse_line(raw):
    line = re.sub(r'^\s*\d+[.、)）]\s*', '', normalize(raw))
    contact = blank_contact()
    contact['notes'] = raw.strip()
    problems = []
    matches = list(LABELS.finditer(line))
    prefix = line[:matches[0].start()] if matches else line
    names = [x for x in re.split(r'[\s,，;；|\t]+', _without_phones(prefix).removesuffix('的')) if x]
    if len(names) == 1 and _valid_name(names[0]):
        contact['name'] = names[0]
    elif names:
        problems.append('无标签内容无法确定学生与家长姓名，请补充标签')
    _append(contact, 'phone', extract_phones(prefix))
    active_role = None
    for index, match in enumerate(matches):
        label = match.group()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(line)
        value = line[match.end():end].strip(' \t:：,，;；|、()（）')
        key = header_field(label)
        role = role_of(label)
        if role:
            active_role = role
        elif re.search(r'学生|孩子|儿童|幼儿|宝宝|学员', label):
            active_role = None
        elif key == 'phone' and active_role:
            key = phone_field(active_role)
        elif key == 'name' and active_role:
            key = active_role
        if not key:
            continue
        if key == 'address':
            _append(contact, key, value)
            active_role = None
        elif key in PHONE_FIELDS:
            _append(contact, key, extract_phones(value))
            if value and not extract_phones(value):
                problems.append(f'{label}不是有效的11位手机号')
        else:
            name = _without_phones(value)
            if name:
                if not _valid_name(name):
                    problems.append(f'{label}内容不明确，请检查姓名与标签')
                elif key == 'name' and contact['name'] and contact['name'] != name:
                    problems.append('同一条记录出现多个学生姓名，请分行')
                else:
                    if re.search(r'外公|外婆|爷爷|奶奶', label):
                        name = f'{label} {name}'
                    if key != 'name' and contact[key] and contact[key] != name:
                        problems.append('同一关系出现多个家长，请改用家长1、家长2')
                    _append(contact, key, name)
            _append(contact, 'phone' if key == 'name' else phone_field(key), extract_phones(value))
    accepted = any(contact[key] for key in FIELDS if key not in {'address', 'notes'}) if matches else bool(contact['name'] and contact['phone'] and len(names) == 1)
    if accepted:
        _check(contact, problems)
    return contact, problems, accepted


def _parse_lines(text):
    result = Recognition()
    records = []
    continuous = False
    for number, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if not line:
            continuous = False
            continue
        if records and continuous and CONTINUATION.match(line):
            records[-1] = (records[-1][0], records[-1][1] + ' ' + line)
        else:
            records.append((number, line))
        continuous = True
    for number, raw in records:
        contact, problems, accepted = _parse_line(raw)
        if accepted:
            result.contacts.append(contact)
        else:
            result.skipped += 1
        if problems or not accepted:
            result.issues.append(f'第{number}行：{"；".join(problems) or "无法识别，请按示例提供学生姓名和家长电话"}')
    return result


def parse_text(text):
    lines = [line for line in text.splitlines() if line.strip()]
    for separator in ('\t', ',', '，', '|'):
        rows = [line.split(separator) for line in lines]
        if any(sum(header_field(cell) is not None for cell in row) >= 2 for row in rows[:30]):
            return parse_rows(rows)
    return _parse_lines(text)


def parse_rows(rows):
    result = Recognition()
    header = -1
    fields = []
    headers = []
    for index, row in enumerate(rows[:30]):
        candidate = [header_field(cell) for cell in row]
        if sum(field is not None for field in candidate) >= 2 and 'name' in candidate:
            header, fields, headers = index, candidate, [normalize(cell) for cell in row]
            if index + 1 < len(rows) and sum(header_field(cell) is not None for cell in rows[index + 1]) >= 2:
                group = ''
                combined = []
                for top, bottom in zip_longest(row, rows[index + 1], fillvalue=''):
                    top, bottom = normalize(top), normalize(bottom)
                    if top:
                        group = top if re.fullmatch(ROLE_WORDS, top) else ''
                    combined.append(group + bottom if group and bottom else bottom or top)
                headers = combined
                fields = [header_field(cell) for cell in combined]
                header = index + 1
            break
    if header < 0:
        return _parse_lines('\n'.join(' '.join(map(normalize, row)) for row in rows))
    result.has_header = True
    for index, row in enumerate(rows[header + 1:], header + 2):
        if not any(normalize(cell) for cell in row):
            continue
        if sum(fields[col] == header_field(cell) for col, cell in enumerate(row) if col < len(fields) and fields[col]) >= 2:
            continue
        contact = blank_contact()
        problems = []
        original = []
        for col, cell in enumerate(row):
            value = normalize(cell)
            if not value:
                continue
            heading = headers[col] if col < len(headers) else f'第{col + 1}列'
            original.append(f'{heading}: {value}')
            key = fields[col] if col < len(fields) else None
            if key in PHONE_FIELDS:
                _append(contact, key, extract_phones(value))
                if not extract_phones(value):
                    problems.append(f'{heading}不是有效的11位手机号')
            elif key == 'address':
                _append(contact, key, value)
            elif key:
                name = _without_phones(value)
                if _valid_name(name):
                    _append(contact, key, name)
                elif name:
                    problems.append(f'{heading}姓名内容不明确')
                _append(contact, 'phone' if key == 'name' else phone_field(key), extract_phones(value))
            elif re.search(r'家长|监护人|姓名|电话|手机', heading) and '关系' not in heading:
                problems.append(f'未支持的列“{heading}”，原文保留在备注')
        for col, cell in enumerate(row):
            heading = re.sub(r'[\s:：()（）_\-]', '', headers[col]) if col < len(headers) else ''
            if not re.fullmatch(r'(?:家长[12一二]?|监护人[12一二]?)?(?:与学生关系|与幼儿关系|与孩子关系|关系|称谓)', heading):
                continue
            source = 'guardian2' if re.search(r'2|二', heading) else 'guardian'
            relation = normalize(cell)
            target = role_of(relation)
            if target in {'father', 'mother'} and not contact[target] and not contact[phone_field(target)]:
                contact[target], contact[phone_field(target)] = contact[source], contact[phone_field(source)]
                contact[source] = contact[phone_field(source)] = ''
            elif relation and contact[source]:
                contact[source] = f'{relation} {contact[source]}'
        contact['notes'] = '；'.join(original)
        if any(contact[key] for key in FIELDS if key not in {'address', 'notes'}):
            _check(contact, problems)
            result.contacts.append(contact)
        else:
            result.skipped += 1
            problems.append('无法识别学生或家长信息')
        if problems:
            result.issues.append(f'第{index}行：{"；".join(dict.fromkeys(problems))}')
    return result


def contact_key(contact):
    return tuple(contact.get(key, '').strip() for key in FIELDS if key != 'notes')


def filter_contacts(contacts, query):
    words = query.lower().split()
    return [contact for contact in contacts if all(word in ' '.join(contact.get(key, '') for key in FIELDS).lower() for word in words)]


def parse_excel(path):
    path = Path(path)
    if path.suffix.lower() not in {'.xlsx', '.xls'} or path.stat().st_size > 2 * 1024 * 1024:
        raise ValueError('请选择不超过 2 MB 的 .xlsx 或 .xls 表格')
    result = Recognition()
    if path.suffix.lower() == '.xlsx':
        with ZipFile(path) as archive:
            check_zip(archive)
            if sum(info.file_size for info in archive.infolist()) > 32 * 1024 * 1024:
                raise ValueError('表格解压后过大，请拆分表格')
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheets = []
            for sheet in book.worksheets:
                if (sheet.max_row or 0) > CONTACT_LIMIT + 32 or (sheet.max_column or 0) > 100:
                    raise ValueError(f'{sheet.title} 行数或列数过多，请拆分表格')
                rows = list(islice(sheet.iter_rows(values_only=True), CONTACT_LIMIT + 33))
                if len(rows) > CONTACT_LIMIT + 32:
                    raise ValueError(f'{sheet.title} 行数过多，请拆分表格')
                sheets.append((sheet.title, len(rows), sheet.max_column, rows))
        finally:
            book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(file_contents=path.read_bytes(), on_demand=True)
        try:
            sheets = []
            for sheet in book.sheets():
                if sheet.nrows > CONTACT_LIMIT + 32 or sheet.ncols > 100:
                    raise ValueError(f'{sheet.name} 行数或列数过多，请拆分表格')
                sheets.append((sheet.name, sheet.nrows, sheet.ncols,
                               [sheet.row_values(row) for row in range(sheet.nrows)]))
        finally:
            book.release_resources()
    for title, count, columns, rows in sheets:
        if count > CONTACT_LIMIT + 32 or columns > 100:
            raise ValueError(f'{title} 行数或列数过多，请拆分表格')
        if not rows:
            continue
        parsed = parse_rows(rows)
        result.sheets += 1
        result.contacts.extend(parsed.contacts)
        if len(result.contacts) > CONTACT_LIMIT:
            raise ValueError(f'本次识别超过 {CONTACT_LIMIT} 条，请拆分表格')
        result.skipped += parsed.skipped
        result.issues.extend(f'{title}：{issue}' for issue in parsed.issues)
        if not parsed.has_header:
            result.no_headers += 1
    return result


class ContactStore:
    def __init__(self, path):
        self.path = Path(path)

    def load(self):
        if not self.path.exists():
            return []
        data = json.loads(self.path.read_text(encoding='utf-8'))
        if (not isinstance(data, list) or len(data) > CONTACT_LIMIT
                or any(not isinstance(item, dict) or not isinstance(item.get('id'), str)
                       or any(not isinstance(item.get(key, ''), str) or
                              (key not in item and key not in OPTIONAL_OLD_FIELDS)
                              for key in FIELDS) for item in data)):
            raise ValueError('通信簿数据格式无效')
        return [{'id': item['id'], **{key: item.get(key, '') for key in FIELDS}} for item in data]

    def save(self, contacts):
        if len(contacts) > CONTACT_LIMIT:
            raise ValueError(f'最多保存 {CONTACT_LIMIT} 条联系人')
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile('w', encoding='utf-8', dir=self.path.parent,
                                             prefix='.contacts-', suffix='.tmp', delete=False) as stream:
                temporary = Path(stream.name)
                os.chmod(temporary, 0o600)
                json.dump(contacts, stream, ensure_ascii=False, indent=2)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, self.path)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
