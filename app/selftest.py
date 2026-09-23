"""End-to-end smoke suite shipped in the executable; synthetic data only."""
from dataclasses import asdict
from pathlib import Path
from threading import Event
import hashlib
import json
import subprocess
import time
import os
import ctypes
from PIL import Image
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pptx import Presentation
from zipfile import ZipFile
from xml.etree import ElementTree
from app.core.contracts import (JobRequest, PhotoOptions, OfficeOptions, OrganizeOptions,
    TemplateOptions, PdfOptions, SheetOptions, MediaOptions)
from app.registry import TOOLS
from app.tools.media import engines


def run_selftest(root: Path):
    started = time.monotonic()
    root.mkdir(parents=True, exist_ok=True)
    (root / 'python-started.json').write_text(json.dumps({'wall_time': time.time()}), encoding='utf-8')
    source_dir = root / '中文输入'
    source_dir.mkdir(exist_ok=True)
    photo = source_dir / '活动照片.jpg'
    Image.effect_noise((800, 600), 80).convert('RGB').save(photo, quality=100)
    doc = Document()
    doc.add_paragraph('离线烟测材料 < & >')
    doc.add_picture(str(photo))
    document = source_dir / '材料.docx'
    doc.save(document)
    presentation = source_dir / '课件.pptx'
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_picture(str(photo), 0, 0)
    slides.save(presentation)
    transparent = source_dir / '透明照片.png'
    Image.new('RGBA', (100, 100), (20, 80, 30, 120)).save(transparent)
    rotated = source_dir / '带方向照片.jpg'
    exif = Image.Exif()
    exif[274] = 6
    Image.new('RGB', (80, 120), 'green').save(rotated, exif=exif)
    sheet = source_dir / '名册.xlsx'
    workbook = Workbook()
    workbook.active.append(['编号', '姓名'])
    workbook.active.append(['001', '测试材料'])
    workbook.save(sheet)
    cached_sheet = source_dir / '缓存公式.xlsx'
    formula_book = Workbook()
    formula_book.active.append(['数量', '说明'])
    formula_book.active.append(['=1+1', '合成样本'])
    formula_book.save(cached_sheet)
    with ZipFile(cached_sheet) as archive:
        members = {info.filename: (info, archive.read(info)) for info in archive.infolist()}
    name = 'xl/worksheets/sheet1.xml'
    xml = ElementTree.fromstring(members[name][1])
    namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    cell = xml.find('.//s:c[@r="A2"]', namespace)
    cell.find('s:v', namespace).text = '2'
    with ZipFile(cached_sheet, 'w') as archive:
        for entry, (info, data) in members.items():
            archive.writestr(info, ElementTree.tostring(xml) if entry == name else data)
    pdf = source_dir / '测试.pdf'
    writer = PdfWriter()
    writer.add_blank_page(300, 400)
    writer.add_blank_page(400, 500)
    writer.write(pdf)
    ffmpeg, _ = engines()
    video = source_dir / '活动视频.avi'
    subprocess.run([ffmpeg, '-v', 'error', '-y', '-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=12',
                    '-f', 'lavfi', '-i', 'sine=frequency=440', '-t', '1', '-c:v', 'rawvideo',
                    '-pix_fmt', 'yuv420p', '-c:a', 'pcm_s16le', str(video)], check=True,
                   stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    cases = [
        ('photo', (photo,), PhotoOptions()),
        ('photo', (transparent, rotated), PhotoOptions(keep_metadata=True)),
        *[('office', (path,), OfficeOptions(optimize, 75))
          for path in (document, presentation) for optimize in (False, True)],
        *[('organize', (photo, document), OrganizeOptions(mode))
          for mode in ('classify', 'rename', 'archive')],
        ('organize', (photo, photo), OrganizeOptions('duplicates')),
        *[('template', (photo,), TemplateOptions(layout, body='教师输入的真实说明', names='测试姓名'))
          for layout in ('photo_docx', 'photo_pptx', 'notice', 'week', 'labels', 'certificate')],
        *[('pdf', (pdf,), PdfOptions(mode, '2,1'))
          for mode in ('merge', 'split', 'extract', 'rotate')],
        ('pdf', (photo, transparent), PdfOptions('images')),
        ('sheet', (sheet, sheet), SheetOptions()),
        ('sheet', (cached_sheet,), SheetOptions(formula_policy='cached')),
        *[('media', (video,), MediaOptions(mode)) for mode in ('video', 'audio', 'wav')],
        ('media', (video,), MediaOptions('video', start=.1, duration=.5)),
        ('media', (video,), MediaOptions('audio', start=.1, duration=.5)),
    ]
    originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in source_dir.iterdir()}
    records = []
    case_records = []
    for index, (tool, inputs, options) in enumerate(cases):
        result = TOOLS[tool][2](JobRequest(tool, inputs, root / f'output-{index}', options),
                              lambda progress: None, Event())
        if not result or any(r.status != 'success' for r in result):
            raise RuntimeError(f'Smoke test failed: {tool}: {result}')
        records.extend(asdict(r) for r in result)
        case_records.append({'tool': tool, 'options': asdict(options), 'status': 'passed'})
    for path, digest in originals.items():
        if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
            raise RuntimeError('An original file was modified')
    # GUI startup from the same frozen application verifies Qt plugins/resources.
    from PySide6.QtWidgets import QApplication
    from app.ui.window import MainWindow
    application = QApplication.instance() or QApplication([])
    window = MainWindow()
    window.show()
    for index in range(7):
        window.nav.setCurrentRow(index)
        application.processEvents()
    window.nav.setCurrentRow(0)
    application.processEvents()
    window.grab().save(str(root / 'ui.png'))
    window.close()
    (root / 'self-test.json').write_text(json.dumps({'status': 'passed', 'cases': len(cases),
        'elapsed_seconds': round(time.monotonic() - started, 2),
        'is_admin': ctypes.windll.shell32.IsUserAnAdmin() != 0 if os.name == 'nt' else os.geteuid() == 0, 'originals_unchanged': True,
        'case_matrix': case_records, 'results': records}, default=str, ensure_ascii=False, indent=2), encoding='utf-8')
