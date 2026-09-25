import math
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt
from app.core.contracts import TemplateOptions, Progress
from app.core.jobs import validate_inputs, success
from app.core.safe_output import SafeOutputWriter, check_cancel
from app.tools.photo import normalized_bytes, validate_direction
from PIL import Image

LAYOUTS = {'notice', 'week', 'labels', 'certificate', 'photo_docx', 'photo_pptx'}


def photo_grid_geometry(options):
    if any(type(value) is not int or not 1 <= value <= 10 for value in (options.rows, options.columns)):
        raise ValueError('照片行数和列数应为 1–10 的整数')
    if options.page_orientation not in {'portrait', 'landscape'}:
        raise ValueError('请选择纸张纵向或横向')
    if options.image_size_mode not in {'auto', 'manual'}:
        raise ValueError('请选择自动适配或手动尺寸')
    page_width, page_height = (21, 29.7) if options.page_orientation == 'portrait' else (29.7, 21)
    # Reserve margins, heading, metadata, a bounded caption and paragraph spacing.
    cell_width = (page_width - 4) / options.columns
    cell_height = (page_height - 10.7) / options.rows
    max_width, max_height = cell_width - .6, cell_height - 1
    if min(max_width, max_height) < .5:
        raise ValueError('行列过密，图片空间不足；请减少行数或列数')
    width, height = max_width, max_height
    if options.image_size_mode == 'manual':
        width, height = options.image_width_cm, options.image_height_cm
        if any(not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0 for v in (width, height)):
            raise ValueError('图片宽高必须为大于 0 的有限数值（厘米）')
        if width > max_width + 1e-8 or height > max_height + 1e-8:
            raise ValueError(f'手动尺寸超过当前版面，每张最多 {max_width:.2f} × {max_height:.2f} 厘米；请减小尺寸或减少行列')
    return page_width, page_height, cell_width, cell_height, width, height


def validate_options(options):
    if not isinstance(options, TemplateOptions) or options.layout not in LAYOUTS:
        raise ValueError('未知材料模板')
    validate_direction(options.rotation, options.orientation)
    if options.layout == 'photo_docx':
        photo_grid_geometry(options)
    for text, limit, label in [(options.title, 40, '标题'), (options.class_name, 30, '班级'),
                                (options.date, 40, '日期'), (options.body, 1500, '正文')]:
        if len(text) > limit or any(ord(c) < 32 and c not in '\n\t' for c in text):
            raise ValueError(f'{label}过长或包含不可用字符，上限 {limit} 字')
    if any('\n' in text or '\r' in text for text in (options.title, options.class_name, options.date)):
        raise ValueError('标题、班级和日期请使用单行文本')
    if options.layout == 'photo_pptx' and (len(options.body) > 100 or len(options.body.splitlines()) > 2):
        raise ValueError('课件说明最多 100 字、2 行，避免文字溢出')
    if options.layout == 'photo_docx' and len(options.body.splitlines()) > 3:
        raise ValueError('照片材料说明最多 3 行，避免文字溢出')
    if not options.title.strip():
        raise ValueError('请填写标题')
    if options.layout in {'notice', 'week'} and not options.body.strip():
        raise ValueError('请填写正文/周计划，不会自动编造内容')
    names = [line.strip() for line in options.names.splitlines() if line.strip()]
    if options.layout in {'labels', 'certificate'} and not names:
        raise ValueError('请填写姓名，每行一位')
    if options.layout.startswith('photo_') and len(options.body) > 180:
        raise ValueError('照片说明最多 180 字，请精简后生成')
    if any(any(ord(c) < 32 for c in name) for name in names):
        raise ValueError('姓名包含不可用控制字符')
    if len(names) > 200 or any(len(name) > 20 for name in names):
        raise ValueError('最多 200 位姓名，每位不超过 20 字')
    return names


def set_slide_text(shape, text, preferred, minimum, wrap=False):
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = Inches(.08)
    frame.margin_top = frame.margin_bottom = Inches(.025)
    frame.word_wrap = wrap
    width = (shape.width - frame.margin_left - frame.margin_right) / 12700
    height = (shape.height - frame.margin_top - frame.margin_bottom) / 12700
    chosen = None
    for size in range(preferred, minimum - 1, -1):
        # Reserve 1.2 em per character, including full-width Chinese glyphs.
        lines = sum(max(1, math.ceil(len(line) * size * 1.2 / width))
                    for line in text.split('\n')) if wrap else 1
        fits_width = wrap or max((len(line) for line in text.split('\n')), default=0) * size * 1.2 <= width
        if fits_width and lines * size * 1.25 <= height:
            chosen = size
            break
    if chosen is None:
        raise ValueError('文字超出固定课件版式，请缩短标题、说明或页脚')
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.name = 'Microsoft YaHei'
        paragraph.font.size = PptPt(chosen)
        paragraph.space_before = paragraph.space_after = PptPt(0)
        paragraph.line_spacing = 1.25


def run(request, emit, cancel):
    options = request.options
    names = validate_options(options)
    photos = options.layout.startswith('photo_')
    validate_inputs(request, allow_empty=not photos)
    if photos and len(request.inputs) > 200:
        raise ValueError('照片材料单次最多 200 张')
    check_cancel(cancel)
    ppt = options.layout == 'photo_pptx'
    extension = '.pptx' if ppt else '.docx'
    if ppt:
        document = Presentation()
        document.slide_width = Inches(13.333)
        document.slide_height = Inches(7.5)
        for index, source in enumerate(request.inputs, 1):
            check_cancel(cancel)
            slide = document.slides.add_slide(document.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(.6), Inches(.3), Inches(12), Inches(.7))
            set_slide_text(heading, options.title, 28, 14)
            data = normalized_bytes(source, rotation=options.rotation, orientation=options.orientation)
            with Image.open(data) as image:
                width, height = image.size
            data.seek(0)
            scale = min(11 / width, (4.8 if options.body else 5.5) / height)
            slide.shapes.add_picture(data, Inches((13.333 - width * scale) / 2), Inches(1.25),
                                     width=Inches(width * scale), height=Inches(height * scale))
            if options.body:
                caption = slide.shapes.add_textbox(Inches(.65), Inches(6.15), Inches(12), Inches(.65))
                set_slide_text(caption, options.body, 14, 11, wrap=True)
            footer = slide.shapes.add_textbox(Inches(.6), Inches(6.9), Inches(12), Inches(.4))
            set_slide_text(footer, f'{options.class_name}  {options.date}  ·  {index}', 12, 8)
            emit(Progress(index, len(request.inputs), '正在生成照片课件'))
    else:
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        if options.layout == 'photo_docx':
            page_width, page_height, cell_width, cell_height, box_width, box_height = photo_grid_geometry(options)
            section.orientation = WD_ORIENT.LANDSCAPE if options.page_orientation == 'landscape' else WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Cm(page_width), Cm(page_height)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2)
        style = document.styles['Normal']
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(11)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        document.add_heading(options.title, 0)
        document.add_paragraph(f'{options.class_name}    {options.date}')
        if options.layout == 'photo_docx':
            if options.body:
                document.add_paragraph(options.body)
            rows, columns = options.rows, options.columns
            per_page = rows * columns
            for offset in range(0, len(request.inputs), per_page):
                check_cancel(cancel)
                if offset:
                    document.add_page_break()
                table = document.add_table(rows=rows, cols=columns)
                table.autofit = False
                for column in table.columns:
                    column.width = Cm(cell_width)
                for row in table.rows:
                    row.height = Cm(cell_height)
                    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell in row.cells:
                        cell.width = Cm(cell_width)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        paragraph = cell.paragraphs[0]
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.paragraph_format.space_before = Pt(0)
                        paragraph.paragraph_format.space_after = Pt(0)
                        paragraph.paragraph_format.line_spacing = 1
                for slot, source in enumerate(request.inputs[offset:offset + per_page]):
                    check_cancel(cancel)
                    data = normalized_bytes(source, rotation=options.rotation, orientation=options.orientation)
                    with Image.open(data) as image:
                        w, h = image.size
                    data.seek(0)
                    if options.image_size_mode == 'auto' or options.keep_aspect_ratio:
                        scale = min(box_width / w, box_height / h)
                        picture_width, picture_height = w * scale, h * scale
                    else:
                        picture_width, picture_height = box_width, box_height
                    cell = table.cell(slot // columns, slot % columns)
                    cell.paragraphs[0].add_run().add_picture(data, width=Cm(picture_width), height=Cm(picture_height))
                    caption = cell.add_paragraph(f'照片 {offset + slot + 1}')
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.paragraph_format.space_before = Pt(0)
                    caption.paragraph_format.space_after = Pt(0)
                    caption.paragraph_format.line_spacing = 1
                    caption.runs[0].font.size = Pt(9)
                    emit(Progress(offset + slot + 1, len(request.inputs), '正在生成 A4 照片材料'))
        elif options.layout == 'labels':
            table = document.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            for index in range(0, len(names), 2):
                check_cancel(cancel)
                cells = table.add_row().cells
                for cell, name in zip(cells, names[index:index + 2]):
                    cell.text = f'\n{name}\n{options.class_name}\n'
        elif options.layout == 'certificate':
            for index, name in enumerate(names):
                check_cancel(cancel)
                if index:
                    document.add_page_break()
                    document.add_heading(options.title, 0)
                document.add_heading(name, 1)
                document.add_paragraph(options.body or '请由教师填写真实的表彰事由。')
                document.add_paragraph(f'{options.class_name}\n{options.date}')
        elif options.layout == 'week':
            table = document.add_table(rows=1, cols=2)
            table.style = 'Table Grid'
            table.rows[0].cells[0].text = '项目'
            table.rows[0].cells[1].text = '计划内容'
            for index, line in enumerate(options.body.splitlines(), 1):
                check_cancel(cancel)
                cells = table.add_row().cells
                parts = line.replace('：', ':').split(':', 1)
                cells[0].text = parts[0] if len(parts) == 2 else str(index)
                cells[1].text = parts[-1]
        else:
            for line in options.body.splitlines():
                document.add_paragraph(line)
    with SafeOutputWriter(request.output_dir / f'材料成品{extension}', request.inputs) as writer:
        document.save(writer.path)
        output = writer.commit(lambda p: Presentation(p) if ppt else Document(p), cancel)
    source = request.inputs[0] if request.inputs else request.output_dir
    return (success(source, output, '固定模板已生成；请复核分页、字体和打印效果'),)
