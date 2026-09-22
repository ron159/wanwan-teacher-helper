from contextlib import ExitStack
from zipfile import ZipFile
from openpyxl import load_workbook, Workbook
from app.core.contracts import SheetOptions, Progress
from app.core.jobs import validate_inputs, success
from app.core.safe_output import SafeOutputWriter, check_cancel
from app.tools.office import check_zip


def run(request, emit, cancel):
    validate_inputs(request)
    options = request.options
    if not isinstance(options, SheetOptions) or options.formula_policy not in {'reject', 'cached'}:
        raise ValueError('表格公式策略无效')
    output_book = Workbook()
    output = output_book.active
    output.title = '汇总'
    expected = None
    rows = 0
    with ExitStack() as stack:
        for index, source in enumerate(request.inputs, 1):
            check_cancel(cancel)
            if source.suffix.lower() != '.xlsx':
                raise ValueError('仅支持无宏 .xlsx 表格')
            with ZipFile(source) as archive:
                check_zip(archive, cancel)
            workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            stack.callback(workbook.close)
            cached = None
            if options.formula_policy == 'cached':
                cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
                stack.callback(cached.close)
            name = options.sheet_name or workbook.sheetnames[0]
            if name not in workbook.sheetnames:
                raise ValueError(f'第 {index} 个文件缺少指定工作表')
            ws = workbook[name]
            if ws.max_row > 100001 or ws.max_column > 200:
                raise ValueError('单表超过 100000 行或 200 列上限')
            iterator = ws.iter_rows()
            header_cells = next(iterator, ())
            header = tuple(cell.value for cell in header_cells)
            if (not header or any(not isinstance(x, str) or not x.strip() for x in header)
                    or len(set(header)) != len(header)):
                raise ValueError('首行必须是非空、不重复的字段名')
            if expected is None:
                expected = header
                output.append([*header, '来源序号'])
            elif header != expected:
                raise ValueError(f'第 {index} 个文件字段或顺序不同，未生成汇总')
            cached_rows = iter(cached[name].iter_rows()) if cached else None
            if cached_rows:
                next(cached_rows)
            for cells in iterator:
                check_cancel(cancel)
                cache_cells = next(cached_rows) if cached_rows else None
                if all(cell.value is None for cell in cells):
                    continue
                values = []
                for col, cell in enumerate(cells):
                    value = cell.value
                    if cell.data_type == 'f':
                        if not cache_cells:
                            raise ValueError(f'第 {index} 个文件含公式；可选择已缓存值策略')
                        value = cache_cells[col].value
                        if value is None:
                            raise ValueError('公式没有缓存值，请先在 Excel/WPS 重新计算并保存')
                    values.append(value)
                output.append([*values, index])
                for col, original in enumerate(cells, 1):
                    target = output.cell(output.max_row, col)
                    target.number_format = original.number_format
                    if isinstance(target.value, str):
                        target.data_type = 's'
                rows += 1
                if rows > 100000:
                    raise ValueError('汇总超过 100000 行上限')
            emit(Progress(index, len(request.inputs), f'已核对 {index} 个表格'))
    output.freeze_panes = 'A2'
    output.auto_filter.ref = output.dimensions
    with SafeOutputWriter(request.output_dir / '同结构汇总.xlsx', request.inputs) as writer:
        output_book.save(writer.path)
        def verify(path):
            book = load_workbook(path, read_only=True)
            try:
                if book.active.max_row != rows + 1:
                    raise ValueError('汇总行数校验失败')
            finally:
                book.close()
        path = writer.commit(verify, cancel)
    output_book.close()
    return (success(request.inputs[0], path, f'已汇总 {rows} 行；编号文本、日期格式保留'),)
