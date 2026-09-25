from threading import Event
import pytest
from PIL import Image
from pypdf import PdfWriter, PdfReader
from openpyxl import Workbook, load_workbook
from docx import Document
from pptx import Presentation
from app.core.contracts import JobRequest, PdfOptions, SheetOptions, TemplateOptions
from app.tools import pdf, sheet, template


def test_pdf_merge_extract_rotate_split(tmp_path):
    source = tmp_path / 'a.pdf'
    writer = PdfWriter()
    for i in range(3):
        writer.add_blank_page(width=300 + i, height=400)
    writer.write(source)
    for options, count in [(PdfOptions('merge'), 3), (PdfOptions('extract', '3,1'), 2),
                            (PdfOptions('rotate', '2', 90), 3)]:
        result = pdf.run(JobRequest('pdf', (source,), tmp_path / 'out', options), lambda p: None, Event())
        pages = PdfReader(result[0].output).pages
        assert len(pages) == count
        if options.mode == 'extract':
            assert pages[0].mediabox.width == 302
        if options.mode == 'rotate':
            assert pages[1].rotation == 90 and pages[0].rotation == 0
    result = pdf.run(JobRequest('pdf', (source,), tmp_path / 'out', PdfOptions('split')),
                     lambda p: None, Event())
    assert len(result) == 3 and all(len(PdfReader(r.output).pages) == 1 for r in result)
    with pytest.raises(ValueError):
        pdf.parse_pages('0,99', 3)


def test_sheet_preserves_ids_and_rejects_mismatch_formula(tmp_path):
    files = (tmp_path / 'a.xlsx', tmp_path / 'b.xlsx')
    for path in files:
        wb = Workbook()
        wb.active.append(['编号', '姓名'])
        wb.active.append(['001', '张同学'])
        wb.save(path)
    req = JobRequest('sheet', files, tmp_path / 'out', SheetOptions())
    result = sheet.run(req, lambda p: None, Event())
    wb = load_workbook(result[0].output)
    assert wb.active.max_row == 3 and wb.active['A2'].value == '001'
    wb.close()
    wb = Workbook()
    wb.active.append(['错误字段'])
    wb.save(files[1])
    with pytest.raises(ValueError):
        sheet.run(req, lambda p: None, Event())
    wb = Workbook()
    wb.active.append(['编号', '姓名'])
    wb.active.append(['=1+1', '姓名'])
    wb.save(files[1])
    with pytest.raises(ValueError):
        sheet.run(req, lambda p: None, Event())


@pytest.mark.parametrize('layout', ['notice', 'week', 'labels', 'certificate', 'photo_docx', 'photo_pptx'])
def test_controlled_templates(tmp_path, layout):
    image = tmp_path / 'photo.jpg'
    Image.new('RGB', (100, 200), 'red').save(image)
    options = TemplateOptions(layout, '春游 & 分享', '小一班', '2026-09-23',
                              '周一：观察树叶\n周二：手工', '张同学\n李同学')
    result = template.run(JobRequest('template', (image,), tmp_path / 'out', options),
                          lambda p: None, Event())
    assert result[0].output.is_file()
    if layout == 'photo_pptx':
        assert len(Presentation(result[0].output).slides) == 1
    else:
        assert Document(result[0].output).paragraphs[0].text == options.title


def test_excel_header_cannot_become_formula(tmp_path):
    source = tmp_path / 'headers.xlsx'
    wb = Workbook()
    wb.active.append(['=1+1', '姓名'])
    wb.active['A1'].data_type = 's'
    wb.active.append(['001', '测试'])
    wb.save(source)
    req = JobRequest('sheet', (source,), tmp_path / 'out', SheetOptions())
    result = sheet.run(req, lambda p: None, Event())
    out = load_workbook(result[0].output)
    assert out.active['A1'].value == '=1+1' and out.active['A1'].data_type == 's'
    out.close()
    wb.active['A1'] = '=1+1'
    wb.save(source)
    with pytest.raises(ValueError, match='首行'):
        sheet.run(req, lambda p: None, Event())


def test_template_rejects_overflowing_caption_and_control_name(tmp_path):
    options = TemplateOptions('photo_docx', body='长' * 181)
    with pytest.raises(ValueError):
        template.validate_options(options)
    with pytest.raises(ValueError):
        template.validate_options(TemplateOptions('labels', names='坏\x01名字'))


def test_pdf_split_cancel_retains_completed_results(tmp_path, monkeypatch):
    source = tmp_path / 'pages.pdf'
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.add_blank_page(100, 100)
    writer.write(source)
    stop = Event()
    original = pdf.save_pdf
    def save_and_cancel(*args):
        path = original(*args)
        stop.set()
        return path
    monkeypatch.setattr(pdf, 'save_pdf', save_and_cancel)
    result = pdf.run(JobRequest('pdf', (source,), tmp_path / 'out', PdfOptions('split')),
                     lambda p: None, stop)
    assert [r.status for r in result] == ['success', 'cancelled']
    assert result[0].output.is_file()


def test_slide_contains_teacher_caption(tmp_path):
    source = tmp_path / 'photo.png'
    Image.new('RGB', (100, 100), 'white').save(source)
    caption = '孩子们一起观察落叶'
    result = template.run(JobRequest('template', (source,), tmp_path / 'out',
        TemplateOptions('photo_pptx', body=caption)), lambda p: None, Event())
    texts = [shape.text for shape in Presentation(result[0].output).slides[0].shapes if shape.has_text_frame]
    assert caption in texts


def test_ppt_long_chinese_title_and_footer_use_bounded_font(tmp_path):
    source = tmp_path / 'photo.jpg'
    Image.new('RGB', (100, 100)).save(source)
    opts = TemplateOptions('photo_pptx', title='长' * 40, class_name='班' * 30, date='日' * 40)
    result = template.run(JobRequest('template', (source,), tmp_path / 'out', opts), lambda _: None, Event())
    slide = Presentation(result[0].output).slides[0]
    text_shapes = [shape for shape in slide.shapes if shape.has_text_frame]
    assert text_shapes[0].text_frame.paragraphs[0].font.size.pt <= 18
    assert text_shapes[-1].text_frame.paragraphs[0].font.size.pt <= 10


@pytest.mark.parametrize('rows,columns,paper', [(1, 1, 'portrait'), (2, 1, 'portrait'), (2, 2, 'landscape'), (3, 2, 'portrait'), (2, 3, 'landscape')])
def test_photo_word_grid_pagination_and_aspect(tmp_path, rows, columns, paper):
    per_page = rows * columns
    from docx.shared import Cm
    from docx.oxml.ns import qn
    import hashlib
    images = []
    for i in range(7):
        path = tmp_path / f'{i}.png'
        Image.new('RGB', (120, 60) if i % 2 else (40, 100), (i * 30, 50, 100)).save(path)
        images.append(path)
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in images]
    options = TemplateOptions(rows=rows, columns=columns, page_orientation=paper, orientation='landscape', body='活动照片')
    result = template.run(JobRequest('template', tuple(images), tmp_path / 'out', options), lambda _: None, Event())
    doc = Document(result[0].output)
    expected_pages = (7 + per_page - 1) // per_page
    assert len(doc.tables) == expected_pages
    assert len(doc.inline_shapes) == 7
    assert len(doc.element.xpath('.//w:br[@w:type="page"]')) == expected_pages - 1
    page_width, page_height, _, _, max_width, max_height = template.photo_grid_geometry(options)
    assert abs(doc.sections[0].page_width.cm - page_width) < .01
    assert abs(doc.sections[0].page_height.cm - page_height) < .01
    for table in doc.tables:
        assert len(table.columns) == columns
        assert len(table.rows) == per_page // columns
        assert all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in table.rows)
    for i, picture in enumerate(doc.inline_shapes):
        assert abs(picture.width / picture.height - (2 if i % 2 else 2.5)) < .001
        assert picture.width <= Cm(max_width)
        assert picture.height <= Cm(max_height)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in images]
    with pytest.raises(ValueError, match='行数和列数'):
        template.validate_options(TemplateOptions(rows=0))


@pytest.mark.parametrize('keep,expected', [(True, (4, 2)), (False, (4, 3))])
def test_manual_photo_dimensions(tmp_path, keep, expected):
    source = tmp_path / 'wide.png'
    Image.new('RGB', (200, 100), 'red').save(source)
    opts = TemplateOptions(rows=2, columns=3, page_orientation='landscape',
        image_size_mode='manual', image_width_cm=4, image_height_cm=3, keep_aspect_ratio=keep)
    result = template.run(JobRequest('template', (source,), tmp_path / 'out', opts), lambda _: None, Event())
    picture = Document(result[0].output).inline_shapes[0]
    assert abs(picture.width.cm - expected[0]) < .001
    assert abs(picture.height.cm - expected[1]) < .001


@pytest.mark.parametrize('values', [dict(rows=11), dict(columns=-1), dict(rows=1.5),
    dict(page_orientation='invalid'), dict(image_size_mode='unknown'),
    dict(rows=10, page_orientation='landscape'),
    dict(image_size_mode='manual', image_width_cm=100),
    dict(image_size_mode='manual', image_height_cm=float('nan')),
    dict(image_size_mode='manual', image_width_cm=0)])
def test_photo_grid_invalid_layout_rejected_before_output(tmp_path, values):
    source = tmp_path / 'photo.png'
    Image.new('RGB', (40, 50)).save(source)
    output = tmp_path / 'out'
    with pytest.raises(ValueError):
        template.run(JobRequest('template', (source,), output, TemplateOptions(**values)), lambda _: None, Event())
    assert not output.exists()
