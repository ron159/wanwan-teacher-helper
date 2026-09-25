from io import BytesIO
from threading import Event
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from PIL import Image
from docx import Document
from app.tools.office import inspect_package, run
from app.core.contracts import JobRequest, OfficeOptions


def test_optimize_changes_only_jpeg_and_keeps_original(tmp_path):
    photo = BytesIO()
    Image.effect_noise((600, 400), 100).convert('RGB').save(photo, 'JPEG', quality=100)
    photo.seek(0)
    document = Document()
    document.add_paragraph('保留 < & > 文字')
    document.add_picture(photo)
    source = tmp_path / '材料.docx'
    document.save(source)
    original = source.read_bytes()
    details = inspect_package(source)
    assert details['categories']['位图']['stored'] > 0
    results = run(JobRequest('office', (source,), tmp_path / 'out', OfficeOptions(True, 75)),
                  lambda p: None, Event())
    assert results[0].status == 'success'
    output = results[0].output
    assert output.stat().st_size < source.stat().st_size
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if not name.endswith('.jpg') and not name.endswith('.jpeg'):
                assert before.read(name) == after.read(name)
    assert source.read_bytes() == original
    assert Document(output).paragraphs[0].text == '保留 < & > 文字'


@pytest.mark.parametrize('member', ['../escape', '/abs', 'a\\b', '_xmlsignatures/a.xml', 'word/vbaProject.bin'])
def test_rejects_unsafe_packages(tmp_path, member):
    source = tmp_path / 'bad.docx'
    with ZipFile(source, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types/>')
        z.writestr(member, 'bad')
    with pytest.raises(ValueError):
        inspect_package(source)


def test_duplicate_and_bomb_rejected(tmp_path):
    source = tmp_path / 'bomb.docx'
    with ZipFile(source, 'w', ZIP_DEFLATED) as z:
        z.writestr('a', b'0' * 2_000_000)
    with pytest.raises(ValueError):
        inspect_package(source)


def test_no_gain_does_not_publish(tmp_path):
    source = tmp_path / 'small.docx'
    Document().save(source)
    result = run(JobRequest('office', (source,), tmp_path / 'out', OfficeOptions(True)),
                 lambda p: None, Event())
    assert result[0].status == 'skipped'
    assert not list((tmp_path / 'out').glob('*')) if (tmp_path / 'out').exists() else True


def test_duplicate_names_symlink_and_external_entity_rejected(tmp_path):
    from zipfile import ZipInfo
    from app.tools.office import check_zip
    from defusedxml.common import DefusedXmlException
    import warnings
    path = tmp_path / 'malicious.docx'
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        with ZipFile(path, 'w') as z:
            z.writestr('same', 'a')
            z.writestr('same', 'b')
    with ZipFile(path) as z, pytest.raises(ValueError):
        check_zip(z)
    info = ZipInfo('link')
    info.create_system = 3
    info.external_attr = 0o120777 << 16
    with ZipFile(path, 'w') as z:
        z.writestr(info, '/etc/passwd')
    with ZipFile(path) as z, pytest.raises(ValueError):
        check_zip(z)
    with ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', '<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><Types>&x;</Types>')
        z.writestr('_rels/.rels', '<Relationships/>')
        z.writestr('word/document.xml', '<document/>')
    with pytest.raises((ValueError, DefusedXmlException)):
        inspect_package(path)


def test_document_jpeg_metadata_is_not_discarded():
    from app.tools.office import optimized_jpeg
    exif = Image.Exif()
    exif[274] = 6
    buffer = BytesIO()
    Image.new('RGB', (100, 100)).save(buffer, 'JPEG', exif=exif, quality=100)
    assert optimized_jpeg(buffer.getvalue(), 70) is None


def test_complex_jpeg_app_metadata_is_preserved_by_skipping():
    from app.tools.office import optimized_jpeg
    buffer = BytesIO()
    Image.effect_noise((100, 100), 80).convert('RGB').save(buffer, 'JPEG', quality=100)
    raw = buffer.getvalue()
    payload = b'Photoshop 3.0\x00IPTC metadata'
    tagged = raw[:2] + b'\xff\xed' + (len(payload) + 2).to_bytes(2, 'big') + payload + raw[2:]
    assert optimized_jpeg(tagged, 70) is None


@pytest.mark.parametrize('member', ['_xmlsignatures/sig1.xml', 'word/vbaProject.bin'])
def test_macro_and_signature_rejection_reaches_specific_guard(tmp_path, member):
    source = tmp_path / 'document.docx'
    Document().save(source)
    with ZipFile(source, 'a') as archive:
        archive.writestr(member, '<Signature/>' if member.endswith('.xml') else b'macro-marker')
    original = source.read_bytes()
    with pytest.raises(ValueError, match='签名或宏'):
        inspect_package(source)
    result = run(JobRequest('office', (source,), tmp_path / 'out', OfficeOptions(True)),
                 lambda p: None, Event())
    assert result[0].status == 'failed'
    assert source.read_bytes() == original
    assert not list((tmp_path / 'out').glob('*'))


def test_ppt_optimization_preserves_reuse_crop_transparency_and_animation(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.oxml.xmlchemy import OxmlElement
    photo = tmp_path / 'photo.jpg'
    Image.effect_noise((600, 400), 100).convert('RGB').save(photo, quality=100)
    transparent = tmp_path / 'transparent.png'
    Image.new('RGBA', (40, 40), (255, 0, 0, 64)).save(transparent)
    animated = tmp_path / 'animated.gif'
    Image.new('RGB', (30, 30), 'red').save(animated, save_all=True,
        append_images=[Image.new('RGB', (30, 30), 'blue')], duration=200, loop=0)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = slide.shapes.add_picture(str(photo), Inches(0), Inches(0), width=Inches(2))
    picture.crop_left = .2
    slide.shapes.add_picture(str(photo), Inches(3), Inches(0), width=Inches(2))
    slide.shapes.add_picture(str(transparent), Inches(0), Inches(3))
    slide.shapes.add_picture(str(animated), Inches(2), Inches(3))
    transition = OxmlElement('p:transition')
    transition.append(OxmlElement('p:fade'))
    slide.element.append(transition)
    source = tmp_path / 'mixed.pptx'
    deck.save(source)
    original = source.read_bytes()
    result = run(JobRequest('office', (source,), tmp_path / 'out', OfficeOptions(True, 70)),
                 lambda _: None, Event())
    assert result[0].status == 'success'
    with ZipFile(source) as before, ZipFile(result[0].output) as after:
        assert before.namelist() == after.namelist()
        jpeg = [n for n in before.namelist() if n.startswith('ppt/media/') and n.endswith(('.jpg', '.jpeg'))]
        assert len(jpeg) == 1  # OOXML deduplicates the twice-used picture.
        for name in before.namelist():
            if name not in jpeg:
                assert before.read(name) == after.read(name), name
        assert after.testzip() is None
    reopened = Presentation(result[0].output)
    assert reopened.slides[0].shapes[0].crop_left == .2
    assert source.read_bytes() == original


def test_aggressive_mode_shrinks_jpeg_and_png_without_changing_other_parts(tmp_path):
    jpeg = tmp_path / 'large.jpg'
    exif = Image.Exif()
    exif[274] = 6
    Image.effect_noise((1800, 1200), 100).convert('RGB').save(jpeg, quality=100, exif=exif)
    png = tmp_path / 'large.png'
    picture = Image.effect_noise((1400, 900), 80).convert('RGBA')
    picture.putalpha(128)
    picture.save(png)
    document = Document()
    document.add_paragraph('原有文字与版式')
    document.add_picture(str(jpeg))
    document.add_picture(str(png))
    source = tmp_path / 'large.docx'
    document.save(source)
    original = source.read_bytes()

    result = run(JobRequest('office', (source,), tmp_path / 'out', OfficeOptions(True, 45, True)),
                 lambda _: None, Event())[0]
    assert result.status == 'success'
    assert result.output.stat().st_size < source.stat().st_size / 2
    assert source.read_bytes() == original
    with ZipFile(source) as before, ZipFile(result.output) as after:
        assert before.namelist() == after.namelist()
        changed = set(result.details['changed'])
        assert changed == {'word/media/image1.jpg', 'word/media/image2.png'}
        for name in before.namelist():
            if name not in changed:
                assert before.read(name) == after.read(name)
        with Image.open(BytesIO(after.read('word/media/image1.jpg'))) as image:
            assert image.format == 'JPEG' and image.size == (853, 1280)
            assert not image.getexif() and not image.info.get('icc_profile')
        with Image.open(BytesIO(after.read('word/media/image2.png'))) as image:
            assert image.format == 'PNG' and max(image.size) <= 1280
            assert image.convert('RGBA').getpixel((0, 0))[3] < 255
    assert Document(result.output).paragraphs[0].text == '原有文字与版式'
