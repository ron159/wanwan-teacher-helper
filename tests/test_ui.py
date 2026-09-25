from PIL import Image
from app.ui.window import MainWindow
from PySide6.QtCore import Qt


def test_default_output_uses_system_documents_location(qtbot, tmp_path, monkeypatch):
    from PySide6.QtCore import QStandardPaths
    redirected = tmp_path / '重定向文档'
    monkeypatch.setattr(QStandardPaths, 'writableLocation', lambda _: str(redirected))
    window = MainWindow()
    qtbot.addWidget(window)
    assert window.output.text() == str(redirected / '丸丸小帮手输出')
    assert not redirected.exists()  # Selecting a default must not create folders.


def test_photo_user_flow_and_all_navigation(qtbot, tmp_path, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    window.show()
    for i in range(7):
        window.nav.setCurrentRow(i)
        assert window.forms.currentIndex() == i
        assert window.title.text()
    window.nav.setCurrentRow(0)
    image = tmp_path / '活动.jpg'
    Image.new('RGB', (100, 100), 'red').save(image)
    window.add_paths([image])
    window.output.setText(str(tmp_path / 'out'))
    monkeypatch.setattr(window, 'confirm_preview', lambda text: True)
    qtbot.mouseClick(window.start_button, Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: bool(window.results) and not window.is_busy(), timeout=10000)
    assert window.results[0].status == 'success'
    assert window.results[0].output.is_file()
    assert window.last_report.is_file()
    assert window.start_button.isEnabled()


def test_queue_order_and_no_execution_without_confirmation(qtbot, tmp_path, monkeypatch):
    window = MainWindow()
    qtbot.addWidget(window)
    paths = [tmp_path / 'a.txt', tmp_path / 'b.txt']
    for path in paths:
        path.write_text('hello')
    window.add_paths(paths)
    window.table.selectRow(1)
    window.move(-1)
    assert window.paths == list(reversed(paths))
    window.nav.setCurrentRow(2)
    window.output.setText(str(tmp_path / 'out'))
    monkeypatch.setattr(window, 'confirm_preview', lambda text: False)
    window.start_preview()
    qtbot.waitUntil(lambda: not window.is_busy(), timeout=10000)
    assert not (tmp_path / 'out').exists()


def test_folder_to_six_photo_word(qtbot, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    from docx import Document
    from app.registry import preview
    from app.core.contracts import JobRequest
    from threading import Event
    window = MainWindow()
    qtbot.addWidget(window)
    window.nav.setCurrentRow(3)
    for i in range(7):
        Image.new('RGB', (100, 50) if i % 2 else (50, 100), 'blue').save(tmp_path / f'{i}.png')
    (tmp_path / 'notes.txt').write_text('not a picture')
    monkeypatch.setattr(QFileDialog, 'getExistingDirectory', lambda *args: str(tmp_path))
    window.add_folder()
    assert len(window.paths) == 7
    form = window.form_widgets['template']
    form.fields['orientation'].setCurrentIndex(1)
    options = form.options()
    assert options.rows == 3 and options.columns == 2 and options.orientation == 'landscape'
    text = preview(JobRequest('template', tuple(window.paths), tmp_path / 'out', options), lambda _: None, Event())
    assert '每页 6 张，共 2 页' in text and '统一横向' in text
    window.output.setText(str(tmp_path / 'out'))
    monkeypatch.setattr(window, 'confirm_preview', lambda _: True)
    window.start_preview()
    qtbot.waitUntil(lambda: bool(window.results) and not window.is_busy(), timeout=10000)
    assert window.results[0].status == 'success'
    doc = Document(window.results[0].output)
    assert len(doc.tables) == 2 and len(doc.inline_shapes) == 7


def test_custom_grid_controls_generate_landscape_manual_word(qtbot, tmp_path, monkeypatch):
    from docx import Document
    window = MainWindow()
    qtbot.addWidget(window)
    window.show()
    window.nav.setCurrentRow(3)
    form = window.form_widgets['template']
    assert form.fields['image_width_cm'].isHidden()
    form.fields['rows'].setValue(2)
    form.fields['columns'].setValue(3)
    form.fields['page_orientation'].setCurrentIndex(1)
    form.fields['image_size_mode'].setCurrentIndex(1)
    assert not form.fields['image_width_cm'].isHidden()
    form.fields['image_width_cm'].setValue(4)
    form.fields['image_height_cm'].setValue(3)
    form.fields['keep_aspect_ratio'].setChecked(False)
    source = tmp_path / 'photo.png'
    Image.new('RGB', (100, 200), 'green').save(source)
    window.add_paths([source])
    window.output.setText(str(tmp_path / 'out'))
    monkeypatch.setattr(window, 'confirm_preview', lambda _: True)
    window.start_preview()
    qtbot.waitUntil(lambda: bool(window.results) and not window.is_busy(), timeout=10000)
    assert window.results[0].status == 'success'
    doc = Document(window.results[0].output)
    assert doc.sections[0].page_width > doc.sections[0].page_height
    assert len(doc.tables[0].rows) == 2 and len(doc.tables[0].columns) == 3
    picture = doc.inline_shapes[0]
    assert abs(picture.width.cm - 4) < .001 and abs(picture.height.cm - 3) < .001
